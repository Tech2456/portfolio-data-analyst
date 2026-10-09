"""Rapport commercial mensuel au format Excel, prêt pour le comité.

Lit les exports du pipeline (site/assets/data/*.json) et produit un
classeur de 5 onglets : synthèse, canaux, alertes, zones, contrôle des
ventes. Mise en forme conditionnelle et graphique natif Excel inclus,
pour que le fichier se lise sans retraitement.

Usage : python -m reporting.rapport_mensuel   ->  reporting/out/rapport_<mois>.xlsx
"""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.formatting.rule import CellIsRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "site" / "assets" / "data"
OUT_DIR = Path(__file__).resolve().parent / "out"

NAVY = "0B2545"
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
HEADER_FONT = Font(bold=True, color="FFFFFF")
RED_FILL = PatternFill("solid", fgColor="F8D7DA")
AMBER_FILL = PatternFill("solid", fgColor="FFF3CD")
GREEN_FILL = PatternFill("solid", fgColor="D4EDDA")
THIN = Border(bottom=Side(style="thin", color="D9D9D9"))

MONTHS_FR = ["janvier", "février", "mars", "avril", "mai", "juin", "juillet",
             "août", "septembre", "octobre", "novembre", "décembre"]


def load(name: str):
    return json.loads((DATA / f"{name}.json").read_text(encoding="utf-8"))


def write_table(ws, top_row: int, headers: list[str], rows: list[list]) -> None:
    for col, title in enumerate(headers, start=1):
        cell = ws.cell(row=top_row, column=col, value=title)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r, row in enumerate(rows, start=top_row + 1):
        for c, value in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=value).border = THIN
    for col in range(1, len(headers) + 1):
        width = max(len(str(headers[col - 1])), *(len(str(r[col - 1])) for r in rows)) if rows else 12
        ws.column_dimensions[get_column_letter(col)].width = min(max(width + 2, 10), 34)
    ws.freeze_panes = ws.cell(row=top_row + 1, column=1)


def month_label(iso: str) -> str:
    year, month = iso[:4], int(iso[5:7])
    return f"{MONTHS_FR[month - 1]} {year}"


def build() -> Path:
    kpis, park = load("kpi_ventes_canal"), load("parc_reabonnement")
    alerts, ruptures = load("alertes"), load("ruptures")
    zones, agents = load("potentiel_zones"), load("ventes_suspectes")

    report_month = park[-1]["month"]
    current = [k for k in kpis if k["month"] == report_month]
    total_sales = sum(k["sales"] for k in current)
    total_target = sum(k["target_sales"] for k in current)
    last, prev = park[-1], park[-2]

    wb = Workbook()

    # 1. Synthèse
    ws = wb.active
    ws.title = "Synthèse"
    ws["A1"] = f"FibreNet Faso — Rapport commercial de {month_label(report_month)}"
    ws["A1"].font = Font(bold=True, size=16, color=NAVY)
    ws["A2"] = "Données simulées — projet portfolio · généré automatiquement par reporting/rapport_mensuel.py"
    ws["A2"].font = Font(italic=True, color="808080")
    tiles = [
        ("Ventes du mois", total_sales, f"objectif {total_target} ({100 * total_sales / total_target:.1f} %)"),
        ("Parc actif", last["active_park"], f"{last['active_park'] - prev['active_park']:+d} vs M-1"),
        ("Taux de réabonnement", last["renewal_rate_pct"] / 100, f"M-1 : {prev['renewal_rate_pct']} %"),
        ("Revenu du mois (FCFA)", last["revenue_fcfa"], f"ARPU {last['arpu_fcfa']:,.0f} FCFA".replace(",", " ")),
    ]
    for i, (label, value, note) in enumerate(tiles):
        col = 1 + i * 2
        ws.cell(row=4, column=col, value=label).font = Font(bold=True, color="595959")
        cell = ws.cell(row=5, column=col, value=value)
        cell.font = Font(bold=True, size=18, color=NAVY)
        cell.number_format = "0.0%" if isinstance(value, float) else "# ##0"
        ws.cell(row=6, column=col, value=note).font = Font(color="808080")
        ws.column_dimensions[get_column_letter(col)].width = 26
    ws["A8"] = "Points d'attention"
    ws["A8"].font = Font(bold=True, size=13, color=NAVY)
    attention = [f"• {a['channel']} à {a['zone']} : {a['deviation_pct']} % vs moyenne 3 mois ({a['severity']})"
                 for a in alerts]
    attention += [f"• Rupture durable (CUSUM) : {r['channel']} à {r['zone']}, part passée de "
                  f"{r['share_before_pct']} % à {r['share_after_pct']} % depuis {r['detected_month']}"
                  for r in ruptures]
    attention += [f"• Contrôle : agent {a['agent_id']} ({a['zone']}), score de risque {a['risk_score']}, "
                  f"{a['never_activated_pct']} % de ventes jamais activées"
                  for a in agents if a["red_flags"] >= 2]
    for i, line in enumerate(attention or ["• Aucun signal anormal ce mois-ci."], start=9):
        ws.cell(row=i, column=1, value=line)

    # 2. Canaux
    ws = wb.create_sheet("Canaux")
    write_table(ws, 1, ["Canal", "Ventes", "Objectif", "Atteinte", "Contribution", "Évol. M-1"],
                [[k["channel"], k["sales"], k["target_sales"], k["attainment_pct"] / 100,
                  k["contribution_pct"] / 100, (k["mom_pct"] or 0) / 100] for k in current])
    for col in "DEF":
        for cell in ws[col][1:]:
            cell.number_format = "0.0%"
    rng = f"D2:D{len(current) + 1}"
    ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["0.9"], fill=RED_FILL))
    ws.conditional_formatting.add(rng, CellIsRule(operator="between", formula=["0.9", "0.99"], fill=AMBER_FILL))
    ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThanOrEqual", formula=["1"], fill=GREEN_FILL))

    # Historique + graphique natif
    channels = sorted({k["channel"] for k in kpis})
    months = sorted({k["month"] for k in kpis})
    by_key = {(k["month"], k["channel"]): k["sales"] for k in kpis}
    top = len(current) + 4
    ws.cell(row=top, column=1, value="Historique des ventes").font = Font(bold=True, color=NAVY)
    ws.cell(row=top + 1, column=1, value="Mois")
    for c, ch in enumerate(channels, start=2):
        ws.cell(row=top + 1, column=c, value=ch)
    for r, m in enumerate(months, start=top + 2):
        ws.cell(row=r, column=1, value=m[:7])
        for c, ch in enumerate(channels, start=2):
            ws.cell(row=r, column=c, value=by_key[(m, ch)])
    chart = LineChart()
    chart.title, chart.y_axis.title, chart.height, chart.width = "Ventes mensuelles par canal", "Ventes", 9, 22
    data = Reference(ws, min_col=2, max_col=1 + len(channels), min_row=top + 1, max_row=top + 1 + len(months))
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=top + 2, max_row=top + 1 + len(months)))
    ws.add_chart(chart, "H2")

    # 3. Alertes
    ws = wb.create_sheet("Alertes")
    write_table(ws, 1, ["Type", "Canal", "Zone", "Constat", "Référence", "Écart", "Gravité"],
                [["Chute mensuelle", a["channel"], a["zone"], a["sales_last_month"], a["baseline_3m"],
                  a["deviation_pct"] / 100, a["severity"]] for a in alerts]
                + [["Rupture CUSUM", r["channel"], r["zone"], r["share_after_pct"] / 100,
                    r["share_before_pct"] / 100, (r["share_after_pct"] - r["share_before_pct"]) / 100,
                    f"depuis {r['detected_month']}"] for r in ruptures])
    for cell in ws["F"][1:]:
        cell.number_format = "0.0%"

    # 4. Zones
    ws = wb.create_sheet("Zones")
    write_table(ws, 1, ["Rang", "Zone", "Foyers raccordables", "Abonnés actifs", "Pénétration",
                        "Réabonnement", "Dynamique 3 mois", "Score priorité"],
                [[z["priority_rank"], z["zone"], z["homes_passed"], z["active_subscribers"],
                  z["penetration_pct"] / 100, z["renewal_pct"] / 100, z["momentum_pct"] / 100,
                  z["priority_score"]] for z in zones])
    for col in "EFG":
        for cell in ws[col][1:]:
            cell.number_format = "0.0%"
    ws.conditional_formatting.add(f"H2:H{len(zones) + 1}",
                                  DataBarRule(start_type="num", start_value=0, end_type="num",
                                              end_value=100, color="2A78D6"))

    # 5. Contrôle des ventes
    ws = wb.create_sheet("Contrôle ventes")
    write_table(ws, 1, ["Agent", "Zone", "Ventes 6 mois", "Jamais activées", "Clients réutilisés",
                        "Ventes fin de mois", "Résiliation précoce", "Score risque", "Signaux rouges"],
                [[a["agent_id"], a["zone"], a["sales"], a["never_activated_pct"] / 100,
                  a["recycled_client_pct"] / 100, a["month_end_pct"] / 100,
                  (a["early_churn_pct"] or 0) / 100, a["risk_score"], a["red_flags"]] for a in agents])
    for col in "DEFG":
        for cell in ws[col][1:]:
            cell.number_format = "0.0%"
    ws.conditional_formatting.add(f"I2:I{len(agents) + 1}",
                                  CellIsRule(operator="greaterThanOrEqual", formula=["2"], fill=RED_FILL))

    OUT_DIR.mkdir(exist_ok=True)
    path = OUT_DIR / f"rapport_{report_month[:7]}.xlsx"
    wb.save(path)
    return path


if __name__ == "__main__":
    print(f"Rapport écrit : {build()}")
